"""説明資料 (PPTX) を作る。画面写真・手順・デモ動画は、操作マニュアルと同じ (manual/content.json)。"""
import json, os, subprocess, tempfile
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

c = json.load(open('manual/content.json', encoding='utf-8'))
OUT = '../docs/manual'
AC, FG, MUTED, HOT, SOFT = RGBColor(0x0A, 0x84, 0xFF), RGBColor(0x14, 0x22, 0x3A), RGBColor(0x5A, 0x6A, 0x85), RGBColor(0xFF, 0x7A, 0x00), RGBColor(0xE8, 0xF2, 0xFF)
FONT = 'Noto Sans JP'

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
blank = prs.slide_layouts[6]


def text(slide, x, y, w, h, s, size=18, bold=False, color=FG, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = s
    r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(size), bold, color, FONT
    return tb


def header(slide, title, sub=None):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.9))
    bar.fill.solid(); bar.fill.fore_color.rgb = AC; bar.line.fill.background()
    text(slide, 0.5, 0.12, 12, 0.7, title, 28, True, RGBColor(255, 255, 255), anchor=MSO_ANCHOR.MIDDLE)
    if sub:
        text(slide, 0.5, 1.0, 12.3, 0.8, sub, 16, False, MUTED)


def badge(slide, x, y, n):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(0.42), Inches(0.42))
    s.fill.solid(); s.fill.fore_color.rgb = HOT; s.line.fill.background()
    tf = s.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = str(n); r.font.size = Pt(16); r.font.bold = True; r.font.color.rgb = RGBColor(255, 255, 255); r.font.name = FONT
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE


def picture(slide, path, x, y, w, h):
    with Image.open(path) as im:
        iw, ih = im.size
    r = min(w / iw, h / ih)
    pw, ph = iw * r, ih * r
    slide.shapes.add_picture(path, Inches(x + (w - pw) / 2), Inches(y + (h - ph) / 2), Inches(pw), Inches(ph))
    return x + (w - pw) / 2, y + (h - ph) / 2, pw, ph


def bullets(slide, x, y, w, items, size=17):
    for i, t in enumerate(items):
        badge(slide, x, y + i * 0.95, i + 1)
        text(slide, x + 0.55, y + i * 0.95 - 0.05, w - 0.55, 0.9, t, size)


# 表紙
s = prs.slides.add_slide(blank)
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height); bg.fill.solid(); bg.fill.fore_color.rgb = AC; bg.line.fill.background()
text(s, 1, 2.2, 11.3, 1.2, c['title'], 44, True, RGBColor(255, 255, 255), PP_ALIGN.CENTER)
text(s, 1, 3.4, 11.3, 0.8, c['subtitle'] + ' 説明資料', 28, False, RGBColor(255, 255, 255), PP_ALIGN.CENTER)
text(s, 2, 4.8, 9.3, 1.4, c['lead'], 16, False, RGBColor(0xE8, 0xF2, 0xFF), PP_ALIGN.CENTER)

# まずはこの4つ
s = prs.slides.add_slide(blank); header(s, 'まずはこの4つ', 'これだけ覚えれば、使い始められます')
for i, st in enumerate(c['start']):
    x, y = 0.6 + (i % 2) * 6.2, 1.9 + (i // 2) * 2.5
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(5.9), Inches(2.2)); box.fill.solid(); box.fill.fore_color.rgb = SOFT; box.line.color.rgb = RGBColor(0xD7, 0xE2, 0xF2)
    text(s, x + 0.25, y + 0.2, 5.4, 0.6, st['t'], 22, True, AC)
    text(s, x + 0.25, y + 0.9, 5.4, 1.2, st['d'], 18)

# 各章
for sec in c['sections']:
    for k, sh in enumerate(sec['shots']):
        s = prs.slides.add_slide(blank); header(s, sec['title'] + ' — ' + sh['cap'])
        top = 1.1
        if k == 0:  # 最初の写真の上に、章の説明を添える
            text(s, 0.5, 1.0, 12.3, 0.7, sec['lead'], 14, False, MUTED)
            top = 1.7
        path = f"manual/out/shots/{sh['file']}.png"
        phone = 'phone' in sh['file']
        if sh['marks']:
            picture(s, path, 0.4, top, 8.3, 7.2 - top)
            bullets(s, 9.0, top + 0.3, 4.0, sh['marks'], 15)
        else:
            picture(s, path, 0.4 if not phone else 4.5, top, 12.5 if not phone else 4.3, 7.2 - top)
    if sec.get('notes'):
        s = prs.slides.add_slide(blank); header(s, sec['title'] + ' — ポイント')
        for i, n in enumerate(sec['notes']):
            text(s, 0.8, 1.4 + i * 1.5, 11.8, 1.4, '● ' + n, 20)

# デモ動画
mp4s = [v for v in c['videos'] if os.path.exists(f"{OUT}/{v['file']}.mp4")]
for v in mp4s:
    s = prs.slides.add_slide(blank); header(s, '操作デモ — ' + v['t'], v['d'])
    mp4 = f"{OUT}/{v['file']}.mp4"
    poster = os.path.join(tempfile.gettempdir(), v['file'] + '.png')
    subprocess.run(['ffmpeg', '-loglevel', 'error', '-y', '-ss', '3', '-i', mp4, '-frames:v', '1', poster], check=True)
    s.shapes.add_movie(mp4, Inches(1.9), Inches(1.9), Inches(9.5), Inches(5.34), poster_frame_image=poster, mime_type='video/mp4')


# 技術の説明 (システム構成・使っている技術・データとセキュリティ・開発と運用)
def node(slide, x, y, w, h, title, sub='', fill=SOFT, tc=AC):
    b = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    b.fill.solid(); b.fill.fore_color.rgb = fill; b.line.color.rgb = RGBColor(0xB8, 0xCC, 0xE8)
    text(slide, x + 0.1, y + 0.08, w - 0.2, 0.45, title, 16, True, tc, PP_ALIGN.CENTER)
    if sub:
        text(slide, x + 0.1, y + 0.55, w - 0.2, h - 0.6, sub, 12, False, FG, PP_ALIGN.CENTER)


def arrow(slide, x1, y1, x2, y2, label=''):
    ln = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    ln.line.color.rgb = MUTED; ln.line.width = Pt(2)
    ln.line._get_or_add_ln().append(ln.line._get_or_add_ln().makeelement('{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd', {'type': 'triangle'}))
    if label:
        text(slide, (x1 + x2) / 2 - 1.0, (y1 + y2) / 2 - 0.32, 2.0, 0.3, label, 11, False, MUTED, PP_ALIGN.CENTER)


s = prs.slides.add_slide(blank); header(s, '技術の説明 — システム構成', 'サーバーを自前で持たず、Googleのサービスと無料枠だけで動いています')
node(s, 0.5, 2.3, 2.4, 1.3, '学生・先生', 'ブラウザ(PC・スマホ)\nで使う')
node(s, 4.0, 1.9, 2.8, 1.0, 'Firebase Hosting', '画面(React)を配信')
node(s, 4.0, 3.2, 2.8, 1.0, 'Firebase Auth', 'Googleログイン\n(学校のアカウントのみ)')
node(s, 4.0, 4.5, 2.8, 1.0, 'Cloud Firestore', 'データの保存\n+ セキュリティルール')
node(s, 8.2, 3.6, 2.4, 1.3, 'Apps Script', '毎日の自動処理\n(バックアップ・進捗書き出し)')
node(s, 10.9, 2.3, 2.1, 1.0, 'ドライブ', 'バックアップ\n進捗報告')
node(s, 10.9, 4.0, 2.1, 1.0, 'Gmail', '先生への通知')
node(s, 0.5, 5.9, 12.4, 1.0, 'GitHub Actions', 'テストして、問題がなければ Hosting とセキュリティルールを公開(手動で実行)', RGBColor(0xFF, 0xF1, 0xE0), HOT)
arrow(s, 2.9, 2.9, 4.0, 2.4); arrow(s, 2.9, 3.0, 4.0, 3.7); arrow(s, 2.9, 3.3, 4.0, 5.0, '直接読み書き')
arrow(s, 8.2, 4.4, 6.8, 5.0, '先生の権限で読む'); arrow(s, 10.6, 4.0, 10.9, 2.9); arrow(s, 10.6, 4.4, 10.9, 4.5)
text(s, 0.5, 5.35, 3.4, 0.5, '※ Cloud Functions・Storage は使わない(無料プランで運用)', 11, False, MUTED)

s = prs.slides.add_slide(blank); header(s, '技術の説明 — 使っている技術', '小さく、無料で、長く動かせることを優先して選んでいます')
rows = [
    ('画面', 'React 18 + TypeScript(厳密モード)+ Vite。UIライブラリは使わず、素のCSS。日本語UI、PC優先でスマホにも対応'),
    ('データ', 'Cloud Firestore(リアルタイムに更新)。サーバーを挟まず、画面から直接読み書きする'),
    ('ログイン', 'Firebase Authentication の Google ログインのみ。学校のメールで、名簿(roster)に載っている人だけが入れる'),
    ('出力', 'Excel(ExcelJS)・PDF(印刷)・CSV。企業に出せる進捗報告書を、画面から作れる'),
    ('自動処理', 'Google Apps Script。毎日のバックアップ、Excel・PDFの進捗書き出し(ドライブへ)、先生への通知(Gmail)'),
    ('公開', 'Firebase Hosting。GitHub Actions で、テストに通ったものだけを公開'),
]
for i, (k, v) in enumerate(rows):
    y = 1.3 + i * 0.95
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(y), Inches(1.9), Inches(0.8)); box.fill.solid(); box.fill.fore_color.rgb = AC; box.line.fill.background()
    text(s, 0.6, y, 1.9, 0.8, k, 18, True, RGBColor(255, 255, 255), PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    text(s, 2.7, y, 10.2, 0.85, v, 15, anchor=MSO_ANCHOR.MIDDLE)

s = prs.slides.add_slide(blank); header(s, '技術の説明 — データとセキュリティ', '守りの本体は、サーバーではなく Firestore のセキュリティルールです')
pts = [
    ('入れる人を限る', '学校のGoogleアカウントで、メールが確認済みで、名簿にある人だけ。名簿にない人は、ログインしてもデータを読めない'),
    ('役割で分ける', '学生は自分のチームだけ。先生は全チームを見られる。キックオフ前の叩き台や名簿など、先生だけが読み書きできるデータがある'),
    ('個人情報を絞る', 'メールアドレスは名簿にだけ置き、ほかのデータは学生ID(A1〜H4)で参照する。他人のメールは読めない'),
    ('消えにくくする', '毎日バックアップをドライブへ。先生は、過去の状態から復元できる'),
    ('試して守る', 'ルールそのものを、エミュレータで自動テストしている。変えるたびに、破られないかを確かめる'),
]
for i, (k, v) in enumerate(pts):
    y = 1.3 + i * 1.12
    badge(s, 0.6, y + 0.1, i + 1)
    text(s, 1.25, y, 3.0, 0.7, k, 19, True, AC)
    text(s, 4.2, y, 8.7, 1.0, v, 15)

s = prs.slides.add_slide(blank); header(s, '技術の説明 — 開発と運用の仕組み', '作ったあとも、壊れずに直し続けられるようにしています')
cols = [
    ('3種類の自動テスト', '・関数の単体テスト(Vitest)\n・セキュリティルールのテスト(Firestoreエミュレータ)\n・画面の通しテスト(Playwright)\n\nこの画面の操作マニュアルと動画も、同じ仕組みで自動で撮影・生成'),
    ('公開の流れ', '1. 変更を保存(GitHub)\n2. 型チェック・テスト・ビルド\n3. すべて通ったら、Hosting とルールを公開\n\n通らなければ、公開されない'),
    ('無料枠で回す工夫', '・読み取り回数を減らす(バックアップは軽い版を毎日、フル版は週1回)\n・Cloud Functions を使わず、定期処理は Apps Script\n・データは小さく保つ(動画や録音は置かない)'),
]
for i, (k, v) in enumerate(cols):
    x = 0.5 + i * 4.2
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.4), Inches(4.0), Inches(5.4)); box.fill.solid(); box.fill.fore_color.rgb = SOFT; box.line.color.rgb = RGBColor(0xD7, 0xE2, 0xF2)
    text(s, x + 0.2, 1.55, 3.6, 0.6, k, 20, True, AC)
    text(s, x + 0.2, 2.3, 3.6, 4.4, v, 14)

# 困ったとき
s = prs.slides.add_slide(blank); header(s, '困ったとき')
for i, f in enumerate(c['faq']):
    y = 1.2 + i * 1.2
    text(s, 0.6, y, 12, 0.4, f['q'], 18, True, AC)
    text(s, 0.6, y + 0.38, 12.1, 0.8, f['a'], 14, False, FG)

prs.save(f"{OUT}/卒研PJ管理_説明資料_学生編.pptx")
print('pptx built')
